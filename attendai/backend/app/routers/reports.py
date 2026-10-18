"""Monthly attendance reports: JSON for the Reports page, plus Excel and PDF downloads."""
import io
from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Annotated, Optional
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Attendance, Department, Employee, Holiday, LeaveRequest
from ..security import require_roles

router = APIRouter(prefix="/api/reports", tags=["Reports"])
IST = ZoneInfo("Asia/Kolkata")
VIEWERS = require_roles("admin", "hr", "manager")
IN_OFFICE = ("present", "late", "half_day")
MonthParam = Annotated[Optional[str], Query(pattern=r"^\d{4}-(0[1-9]|1[0-2])$", description="YYYY-MM")]
STATUS_LABEL = {"present": "On time", "late": "Late", "half_day": "Half day", "absent": "Absent",
                "on_leave": "On leave", "holiday": "Holiday", "weekend": "Weekend"}
MODE_LABEL = {"office": "Office", "wfh": "Work from home", "hybrid": "Hybrid"}


def month_bounds(month: Optional[str]):
    today = datetime.now(IST).date()
    year, mon = map(int, month.split("-")) if month else (today.year, today.month)
    return date(year, mon, 1), date(year, mon, monthrange(year, mon)[1]), today


def build_report(db: Session, user: Employee, month: Optional[str], department_id: Optional[int]) -> dict:
    """One row per employee. Only completed working days count, so today is never marked absent."""
    first, last, today = month_bounds(month)
    stmt = select(Employee).where(Employee.status != "inactive").order_by(Employee.full_name)
    if user.role == "manager":
        stmt = stmt.where(Employee.manager_id == user.id)
    if department_id:
        stmt = stmt.where(Employee.department_id == department_id)
    people = db.scalars(stmt).all()
    ids = [p.id for p in people]

    holidays = set(db.scalars(select(Holiday.holiday_date).where(
        Holiday.holiday_date.between(first, last), Holiday.is_optional.is_(False))))
    through = min(last, today - timedelta(days=1))
    span = [first + timedelta(i) for i in range((through - first).days + 1)] if through >= first else []
    workdays = [d for d in span if d.weekday() < 5 and d not in holidays]

    attendance: dict[int, dict[date, Attendance]] = defaultdict(dict)
    leaves: dict[int, list[LeaveRequest]] = defaultdict(list)
    if ids:
        for a in db.scalars(select(Attendance).where(
                Attendance.employee_id.in_(ids), Attendance.work_date.between(first, last))):
            attendance[a.employee_id][a.work_date] = a
        for req in db.scalars(select(LeaveRequest).where(
                LeaveRequest.employee_id.in_(ids), LeaveRequest.status == "approved",
                LeaveRequest.start_date <= last, LeaveRequest.end_date >= first)):
            leaves[req.employee_id].append(req)

    rows = []
    for p in people:
        mine = attendance[p.id]
        days = {d for d in workdays if d >= p.date_of_joining}
        attended = {d for d, a in mine.items() if a.status in IN_OFFICE and d in days}
        leave_by_day: dict[date, float] = {}
        for req in leaves[p.id]:
            for i in range((req.end_date - req.start_date).days + 1):
                d = req.start_date + timedelta(i)
                if d in days and d not in attended:
                    leave_by_day[d] = 0.5 if req.is_half_day else 1.0
        leave_days = sum(leave_by_day.values())
        minutes = sum(a.worked_minutes for a in mine.values())
        worked_days = sum(1 for a in mine.values() if a.worked_minutes > 0)
        expected = len(days) - leave_days
        rows.append({
            "employee_id": p.id, "employee_code": p.employee_code, "full_name": p.full_name,
            "department": p.department.name if p.department else None,
            "working_days": len(days), "present": len(attended),
            "late": sum(1 for a in mine.values() if a.status == "late"),
            "half_days": sum(1 for a in mine.values() if a.status == "half_day"),
            "leave": leave_days, "absent": max(0.0, len(days) - len(attended) - leave_days),
            "hours": round(minutes / 60, 1),
            "avg_hours": round(minutes / 60 / worked_days, 1) if worked_days else 0.0,
            "overtime_hours": round(sum(a.overtime_minutes for a in mine.values()) / 60, 1),
            "attendance_pct": round(100 * len(attended) / expected, 1) if expected > 0 else None,
        })

    present_total = sum(r["present"] for r in rows)
    expected_total = sum(max(0.0, r["working_days"] - r["leave"]) for r in rows)
    return {
        "month": f"{first.year}-{first.month:02d}", "working_days": len(workdays), "rows": rows,
        "totals": {
            "late": sum(r["late"] for r in rows), "absent": round(sum(r["absent"] for r in rows), 1),
            "leave": round(sum(r["leave"] for r in rows), 1),
            "overtime_hours": round(sum(r["overtime_hours"] for r in rows), 1),
            "attendance_pct": round(100 * present_total / expected_total, 1) if expected_total else None,
        },
    }


def make_xlsx(db: Session, report: dict, team: str, label: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_font, head_fill = Font(name="Arial", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="14213D")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Employee ID", "Name", "Team", "Working days", "Days present", "Late arrivals", "Half days", "Leave days",
               "Absent days", "Total hours", "Avg hours / day", "Overtime hours", "Attendance %"])
    for r in report["rows"]:
        n = ws.max_row + 1
        ws.append([r["employee_code"], r["full_name"], r["department"] or "-", r["working_days"], r["present"], r["late"],
                   r["half_days"], r["leave"], f"=MAX(0,D{n}-E{n}-H{n})", r["hours"], r["avg_hours"], r["overtime_hours"],
                   f'=IFERROR(E{n}/(D{n}-H{n}),"")'])
    last = ws.max_row
    if last > 1:
        t = last + 1
        ws.append(["Total", "", ""] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{last})" for c in range(4, 11)]
                  + [None, f"=SUM(L2:L{last})", f'=IFERROR(E{t}/(D{t}-H{t}),"")'])
        ws.cell(t + 2, 1, "Absent days and attendance % are formulas. See the Notes sheet for how every column is defined.")
        ws.auto_filter.ref = f"A1:M{last}"
    for row in ws.iter_rows():
        for c in row:
            c.font = base
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    if last > 1:
        for c in ws[last + 1]:
            c.font, c.fill = bold, PatternFill("solid", fgColor="E5E9F0")
    for r in range(2, ws.max_row + 1):
        for col in (8, 9, 10, 11, 12):
            ws.cell(r, col).number_format = "0.0"
        ws.cell(r, 13).number_format = "0.0%"
        for col in range(4, 14):
            ws.cell(r, col).alignment = Alignment(horizontal="center")
    for i, w in enumerate([12, 24, 20, 13, 13, 13, 10, 11, 12, 12, 14, 14, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"

    log = wb.create_sheet("Daily log")
    log.append(["Date", "Employee ID", "Name", "Team", "Check in", "Check out", "Worked (hours)", "Late (min)",
                "Overtime (min)", "Status", "Works from"])
    first, last_day, _ = month_bounds(report["month"])
    ids = [r["employee_id"] for r in report["rows"]]
    stmt = (select(Attendance, Employee).join(Employee, Employee.id == Attendance.employee_id)
            .where(Attendance.employee_id.in_(ids), Attendance.work_date.between(first, last_day))
            .order_by(Attendance.work_date, Employee.full_name))
    for a, e in db.execute(stmt):
        log.append([a.work_date, e.employee_code, e.full_name, e.department.name if e.department else "-",
                    a.check_in.astimezone(IST).time() if a.check_in else None,
                    a.check_out.astimezone(IST).time() if a.check_out else None,
                    round(a.worked_minutes / 60, 2), a.late_minutes, a.overtime_minutes,
                    STATUS_LABEL.get(a.status, a.status), MODE_LABEL.get(a.work_mode, a.work_mode)])
    for row in log.iter_rows():
        for c in row:
            c.font = base
    for c in log[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(2, log.max_row + 1):
        log.cell(r, 1).number_format = "dd mmm yyyy"
        log.cell(r, 5).number_format = log.cell(r, 6).number_format = "h:mm AM/PM"
    for i, w in enumerate([13, 12, 24, 20, 11, 11, 14, 11, 14, 11, 16], start=1):
        log.column_dimensions[get_column_letter(i)].width = w
    log.freeze_panes = "A2"
    if log.max_row > 1:
        log.auto_filter.ref = f"A1:K{log.max_row}"

    notes = wb.create_sheet("Notes")
    for line in [
        "AttendAI monthly attendance report", f"Month: {label}. Team: {team}.", "",
        "Working days: Monday to Friday, excluding fixed company holidays, from the employee's joining date to yesterday.",
        "Today is left out until the day ends, so nobody is marked absent for a day still in progress.",
        "Days present: working days with a check-in (on time, late or half day).",
        "Leave days: approved leave on working days when the employee did not check in. A half-day leave counts as 0.5.",
        "Absent days = working days - days present - leave days (formula).",
        "Attendance % = days present / (working days - leave days) (formula).",
        "Times are shown in India Standard Time (IST).",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    for row in notes.iter_rows():
        for c in row:
            c.font = base
    notes["A1"].font = Font(name="Arial", size=12, bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pdf(report: dict, team: str, label: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    navy, grey = colors.HexColor("#14213D"), colors.HexColor("#566377")
    page = landscape(A4)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm,
                            bottomMargin=16 * mm, title=f"AttendAI attendance report {label}")
    styles = getSampleStyleSheet()
    title = ParagraphStyle("title", parent=styles["Title"], alignment=0, fontSize=20, textColor=navy, spaceAfter=2)
    meta = ParagraphStyle("meta", parent=styles["Normal"], fontSize=9, textColor=grey)
    story = [Paragraph("Attendance report", title),
             Paragraph(f"{label}. {team}. {report['working_days']} completed working days.", meta), Spacer(1, 6 * mm)]

    rows = report["rows"]
    if not rows:
        story.append(Paragraph("No employees match this report.", styles["Normal"]))
    else:
        num = lambda v: f"{v:g}"
        pct = lambda v: "-" if v is None else f"{v:.0f}%"
        data = [["Employee", "Team", "Working\ndays", "Present", "Late", "Leave", "Absent", "Hours", "Avg / day", "Overtime", "Attendance"]]
        for r in rows:
            data.append([f"{r['full_name']} ({r['employee_code']})", r["department"] or "-", r["working_days"], r["present"],
                         r["late"], num(r["leave"]), num(r["absent"]), num(r["hours"]), num(r["avg_hours"]),
                         num(r["overtime_hours"]), pct(r["attendance_pct"])])
        t = report["totals"]
        data.append(["Total", "", sum(r["working_days"] for r in rows), sum(r["present"] for r in rows), t["late"],
                     num(t["leave"]), num(t["absent"]), num(round(sum(r["hours"] for r in rows), 1)), "",
                     num(t["overtime_hours"]), pct(t["attendance_pct"])])
        table = Table(data, repeatRows=1, colWidths=[w * mm for w in (58, 40, 18, 18, 16, 16, 16, 18, 18, 20, 23)])
        style = [
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("ALIGN", (2, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F3F5F9")]),
            ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#D6DCE6")),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E5E9F0")), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
        for i, r in enumerate(rows, start=1):
            if r["attendance_pct"] is not None and r["attendance_pct"] < 80:
                style.append(("TEXTCOLOR", (10, i), (10, i), colors.HexColor("#C92F45")))
        table.setStyle(TableStyle(style))
        story.append(table)

    stamp = datetime.now(IST).strftime("%d %b %Y, %I:%M %p IST")

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(grey)
        canvas.drawString(14 * mm, 9 * mm, f"AttendAI, generated {stamp}")
        canvas.drawRightString(page[0] - 14 * mm, 9 * mm, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


@router.get("/monthly")
def monthly(month: MonthParam = None, department_id: Optional[int] = None,
            db: Session = Depends(get_db), user: Employee = Depends(VIEWERS)):
    return build_report(db, user, month, department_id)


@router.get("/monthly/export")
def export_monthly(month: MonthParam = None, department_id: Optional[int] = None,
                   file_format: str = Query("xlsx", alias="format", pattern="^(xlsx|pdf)$"),
                   db: Session = Depends(get_db), user: Employee = Depends(VIEWERS)):
    report = build_report(db, user, month, department_id)
    first, _, _ = month_bounds(report["month"])
    label = first.strftime("%B %Y")
    dept = db.get(Department, department_id) if department_id else None
    if department_id and not dept:
        raise HTTPException(404, "Team not found")
    team = dept.name if dept else "All teams"
    if file_format == "pdf":
        body, mime = make_pdf(report, team, label), "application/pdf"
    else:
        body, mime = make_xlsx(db, report, team, label), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return Response(body, media_type=mime,
                    headers={"Content-Disposition": f'attachment; filename="AttendAI-attendance-{report["month"]}.{file_format}"'})
